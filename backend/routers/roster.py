import json
import os
import re
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

import httpx
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from lib.dates import today_iso
from models.roster import (
    ClassSlot,
    ExportRequest,
    ExtractionResponse,
    MemberSchedule,
    PlotAssignment,
    PlotRequest,
    PlotResponse,
    ValidationItem,
)


router = APIRouter(prefix="/roster", tags=["roster"])
FILE_PATTERN = re.compile(r"\.(?P<ext>[^.]+)$", re.IGNORECASE)
REQUIRED_DIVISIONS = ["RP", "FG", "VG", "CW", "IL", "WM", "PK", "DG"]
SUPPORTED_TEXT_EXTENSIONS = {"txt", "md", "csv", "json"}
MAX_FILE_BYTES = 12 * 1024 * 1024


def _minutes(value: str) -> int:
    parsed = datetime.strptime(value, "%H:%M")
    return parsed.hour * 60 + parsed.minute


def _clock(total_minutes: int) -> str:
    return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"


def _overlaps(start: int, end: int, class_slot: ClassSlot) -> bool:
    return start < _minutes(class_slot.end_time) and end > _minutes(class_slot.start_time)


def _extract_ai_members(raw: str, filename: str) -> list[MemberSchedule]:
    start = raw.find("{")
    end = raw.rfind("}")
    if start < 0 or end < start:
        raise ValueError("respons AI tidak berisi JSON")
    try:
        data = json.loads(raw[start : end + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"respons AI terpotong atau bukan JSON valid pada baris {exc.lineno}, kolom {exc.colno}"
        ) from exc
    members: list[MemberSchedule] = []
    for item in data.get("members", []):
        classes = [ClassSlot(**slot) for slot in item.get("classes", [])]
        if not classes:
            continue
        members.append(
            MemberSchedule(
                code=str(item.get("code", item.get("name", "Anggota"))),
                division=str(item.get("division", "-")),
                generation=str(item.get("generation", "00")),
                source_file=filename,
                confidence=float(item.get("confidence", 0.9)),
                notes=[str(note) for note in item.get("notes", [])],
                classes=classes,
            )
        )
    return members


_DAY_ALIASES = {
    "senin": "Senin",
    "selasa": "Selasa",
    "rabu": "Rabu",
    "kamis": "Kamis",
    "jumat": "Jumat",
    "sabtu": "Sabtu",
    "monday": "Senin",
    "tuesday": "Selasa",
    "wednesday": "Rabu",
    "thursday": "Kamis",
    "friday": "Jumat",
    "saturday": "Sabtu",
}
_TIME_PATTERN = re.compile(r"\b([01]?\d|2[0-3])\s*[:.]\s*([0-5]\d)\b")


def _member_from_text(text: str, code: str, division: str, generation: str, filename: str) -> MemberSchedule:
    classes: list[ClassSlot] = []
    for line in text.splitlines():
        normalized = re.sub(r"\s+", " ", re.sub(r"[#|,;]", " ", line)).strip()
        day_match = re.search(
            r"\b(Senin|Selasa|Rabu|Kamis|Jumat|Sabtu|Monday|Tuesday|Wednesday|Thursday|Friday|Saturday)\b",
            normalized,
            re.IGNORECASE,
        )
        times = list(_TIME_PATTERN.finditer(normalized))
        if not day_match or len(times) < 2:
            continue
        start_time = f"{int(times[0].group(1)):02d}:{times[0].group(2)}"
        end_time = f"{int(times[1].group(1)):02d}:{times[1].group(2)}"
        if start_time >= end_time:
            continue
        course = normalized
        for token in (day_match.group(0), times[0].group(0), times[1].group(0)):
            course = course.replace(token, "", 1)
        course = re.sub(r"[-–—]", " ", course)
        course = re.sub(r"\s+", " ", course).strip() or "Mata kuliah dari file teks"
        classes.append(
            ClassSlot(
                day=_DAY_ALIASES[day_match.group(1).lower()],
                start_time=start_time,
                end_time=end_time,
                course=course,
            )
        )
    unique_classes = list({
        (slot.day, slot.start_time, slot.end_time, slot.course): slot for slot in classes
    }.values())
    return MemberSchedule(
        code=code,
        division=division,
        generation=generation,
        source_file=filename,
        confidence=0.85 if unique_classes else 0,
        notes=[] if unique_classes else ["File teks tidak memiliki baris jadwal yang dapat dikenali."],
        classes=unique_classes,
    )


async def _extract_schedule(
    path: str, mime_type: str, code: str, division: str, generation: str, filename: str
) -> list[MemberSchedule]:
    with open(path, "rb") as f:
        text = f.read().decode("utf-8-sig")
    if filename.lower().endswith(".json"):
        value = json.loads(text)
        items = value if isinstance(value, list) else value.get("classes", [])
        text = "\n".join(
            f"{item.get('day', '')} | {item.get('start_time', '')} | {item.get('end_time', '')} | {item.get('course', '')}"
            for item in items if isinstance(item, dict)
        )
    api_key = os.environ.get("EMERGENT_LLM_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="Kunci layanan AI belum dikonfigurasi")
    if len(text) > 300_000:
        raise HTTPException(status_code=413, detail="File teks terlalu besar untuk diproses AI")

    prompt = f"""Anda adalah parser jadwal kuliah. Baca isi file teks berikut, termasuk tabel Markdown, CSV, atau JSON.
Pisahkan setiap jadwal berdasarkan identitas yang muncul di dalam isi, misalnya heading `RSY_DG_21.jpg`, `Kelas: ...`, `Kode: ...`, atau blok dokumen bernomor.
Jangan gunakan nama file upload sebagai identitas anggota. Jika kode/divisi/angkatan hanya ada pada heading nama berkas di dalam teks, ekstrak dari sana.
Kembalikan JSON valid saja dengan struktur persis:
{{"members":[{{"code":"RSY","division":"DG","generation":"21","confidence":0.95,"notes":[],"classes":[{{"day":"Senin","start_time":"10:30","end_time":"12:10","course":"Pancasila"}}]}}]}}
Hari harus salah satu: Senin, Selasa, Rabu, Kamis, Jumat, Sabtu. Waktu harus HH:MM. Abaikan baris header tabel, catatan, citation, dosen, dan ruang. Jangan membuat jadwal yang tidak ada.

ISI FILE:
{text}"""
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.0,
            "maxOutputTokens": 32768,
            "responseMimeType": "application/json",
            "responseSchema": {
                "type": "OBJECT",
                "properties": {
                    "members": {
                        "type": "ARRAY",
                        "items": {
                            "type": "OBJECT",
                            "properties": {
                                "code": {"type": "STRING"},
                                "division": {"type": "STRING"},
                                "generation": {"type": "STRING"},
                                "confidence": {"type": "NUMBER"},
                                "notes": {"type": "ARRAY", "items": {"type": "STRING"}},
                                "classes": {
                                    "type": "ARRAY",
                                    "items": {
                                        "type": "OBJECT",
                                        "properties": {
                                            "day": {"type": "STRING"},
                                            "start_time": {"type": "STRING"},
                                            "end_time": {"type": "STRING"},
                                            "course": {"type": "STRING"},
                                        },
                                        "required": ["day", "start_time", "end_time", "course"],
                                    },
                                },
                            },
                            "required": ["code", "division", "generation", "classes"],
                        },
                    },
                },
                "required": ["members"],
            },
        },
    }
    model = os.environ.get("GEMINI_TEXT_MODEL", "gemini-3.5-flash-lite")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    try:
        async with httpx.AsyncClient(timeout=55.0) as client:
            response = await client.post(url, headers={"x-goog-api-key": api_key}, json=payload)
            response.raise_for_status()
        response_data = response.json()
        candidate = response_data.get("candidates", [{}])[0]
        if candidate.get("finishReason") == "MAX_TOKENS":
            raise HTTPException(
                status_code=502,
                detail="Respons AI terpotong karena terlalu banyak jadwal. Pecah file menjadi beberapa bagian lalu unggah ulang.",
            )
        parts = candidate.get("content", {}).get("parts", [])
        members = _extract_ai_members("".join(str(part.get("text", "")) for part in parts), filename)
        if not members:
            raise HTTPException(status_code=422, detail="AI tidak menemukan jadwal dalam file teks")
        return members
    except HTTPException:
        raise
    except httpx.HTTPStatusError as exc:
        raise HTTPException(
            status_code=502,
            detail=f"AI text parsing error: {exc.response.status_code} - {exc.response.text}",
        ) from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Pembacaan file teks gagal: {exc}") from exc


@router.post("/extract", response_model=ExtractionResponse)
async def extract_documents(files: list[UploadFile] = File(...)):
    if not files:
        raise HTTPException(status_code=400, detail="Pilih minimal satu file")
    if len(files) > 20:
        raise HTTPException(status_code=400, detail="Maksimal 20 file per proses")

    members: list[MemberSchedule] = []
    warnings: list[str] = []
    for upload in files:
        filename = upload.filename or ""
        match = FILE_PATTERN.search(filename)
        if not match or match.group("ext").lower() not in SUPPORTED_TEXT_EXTENSIONS:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"Nama file {filename or '(tanpa nama)'} tidak valid. Gunakan pola "
                    "Gunakan file teks TXT, MD, CSV, atau JSON. Identitas dibaca dari isi file."
                ),
            )
        suffix = f".{match.group('ext').lower()}"
        with tempfile.TemporaryDirectory(prefix="autoplot-") as temp_dir:
            temp_path = str(Path(temp_dir) / f"source{suffix}")
            total_bytes = 0
            with open(temp_path, "wb") as temp_file:
                while chunk := await upload.read(1024 * 1024):
                    total_bytes += len(chunk)
                    if total_bytes > MAX_FILE_BYTES:
                        raise HTTPException(status_code=413, detail=f"{filename} melebihi batas 12 MB")
                    temp_file.write(chunk)
            extracted_members = await _extract_schedule(
                temp_path,
                upload.content_type,
                "",
                "",
                "",
                filename,
            )
            members.extend(extracted_members)
            if any(member.confidence < 0.7 for member in extracted_members):
                warnings.append(f"Periksa ulang hasil AI dari {filename}: confidence rendah")
    return ExtractionResponse(members=members, warnings=warnings)


@router.post("/plot", response_model=PlotResponse)
async def generate_plot(payload: PlotRequest):
    divisions = sorted({schedule.division for schedule in payload.schedules})
    missing_divisions = [division for division in REQUIRED_DIVISIONS if division not in divisions]
    usage = {schedule.id: 0 for schedule in payload.schedules}
    assignments: list[PlotAssignment] = []
    validations: list[ValidationItem] = []
    coverage_by_day: dict[str, list[str]] = {}
    complete_days = []
    duration = payload.config.shift_duration_hours * 60
    opening = _minutes(payload.config.operating_start)
    closing = _minutes(payload.config.operating_end)

    if missing_divisions:
        validations.append(
            ValidationItem(
                level="error",
                message=f"Batch belum lengkap. Divisi yang belum tersedia: {', '.join(missing_divisions)}.",
            )
        )

    for day in payload.config.active_days:
        if missing_divisions:
            coverage_by_day[day] = []
            continue
        proposed: list[tuple[MemberSchedule, int]] = []
        for division in REQUIRED_DIVISIONS:
            candidates = sorted(
                [member for member in payload.schedules if member.division == division],
                key=lambda member: (usage[member.id], member.code),
            )
            selected: tuple[MemberSchedule, int] | None = None
            for member in candidates:
                day_classes = [slot for slot in member.classes if slot.day == day]
                for start in range(opening, closing - duration + 1, 60):
                    end = start + duration
                    if not any(_overlaps(start, end, slot) for slot in day_classes):
                        selected = (member, start)
                        break
                if selected:
                    break
            if not selected:
                validations.append(
                    ValidationItem(level="error", day=day, message=f"Tidak ada slot bebas untuk divisi {division}; hari tidak diplot.")
                )
                proposed = []
                break
            proposed.append(selected)

        if not proposed:
            coverage_by_day[day] = []
            continue
        for member, start in proposed:
            assignments.append(
                PlotAssignment(
                    day=day,
                    start_time=_clock(start),
                    end_time=_clock(start + duration),
                    member_code=member.code,
                    division=member.division,
                    generation=member.generation,
                )
            )
            usage[member.id] += 1
        coverage_by_day[day] = sorted(member.division for member, _ in proposed)
        complete_days.append(day)

    if complete_days:
        validations.append(
            ValidationItem(level="ok", message=f"{len(complete_days)} hari memenuhi delapan divisi tanpa bentrok jadwal kuliah.")
        )
    else:
        validations.append(ValidationItem(level="error", message="Tidak ada hari yang memenuhi seluruh aturan plotting."))
    return PlotResponse(
        assignments=assignments,
        validations=validations,
        divisions=divisions,
        required_divisions=REQUIRED_DIVISIONS,
        missing_divisions=missing_divisions,
        coverage_by_day=coverage_by_day,
        complete_days=complete_days,
        export_ready=bool(complete_days) and not missing_divisions,
    )


@router.post("/export")
async def export_plot(payload: ExportRequest):
    if not payload.plot.export_ready or not payload.plot.complete_days:
        raise HTTPException(status_code=409, detail="Plot belum memenuhi syarat ekspor delapan divisi")
    workbook = Workbook()
    main = workbook.active
    main.title = "Plotting Utama"
    main.append(["ID Batch", "Hari", "Mulai", "Selesai", "Kode Anggota", "Divisi", "Angkatan"])
    for assignment in payload.plot.assignments:
        main.append([
            payload.batch_id,
            assignment.day,
            assignment.start_time,
            assignment.end_time,
            assignment.member_code,
            assignment.division,
            assignment.generation,
        ])

    validation = workbook.create_sheet("Validasi & Error")
    validation.append(["ID Batch", "Status", "Hari", "Pesan"])
    for item in payload.plot.validations:
        validation.append([payload.batch_id, item.level.upper(), item.day or "-", item.message])

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (main, validation):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 3, 72)
            sheet.column_dimensions[column[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"plot-absensi-{payload.batch_id}-{today_iso()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )